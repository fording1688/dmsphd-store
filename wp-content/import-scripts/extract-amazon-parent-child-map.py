#!/usr/bin/env python3
import csv
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook


BASE_COLS = {
    "status": 1,
    "title": 2,
    "sku": 3,
    "parentage": 6,
    "parent_sku": 7,
    "variation_theme": 8,
    "item_name": 9,
    "brand": 11,
}

ATTR_COLS = {
    "Style": 45,
    "Material": 46,
    "Number of Items": 51,
    "Package Quantity": 52,
    "Wheel Diameter": 53,
    "Color": 55,
    "Size": 56,
    "Shape": 58,
    "Grit Type": 136,
    "Item Diameter": 143,
    "Grit Material": 181,
    "Grit Number": 182,
    "Thread Size": 184,
    "Point Style": 213,
    "Blade Type": 236,
}

THEME_ATTRS = {
    "COLOR": ["Color"],
    "SIZE": ["Size"],
    "COLOR/SIZE": ["Color", "Size"],
    "NUMBER_OF_ITEMS": ["Number of Items"],
    "SIZE/STYLE": ["Size", "Style"],
    "STYLE": ["Style"],
}


def cell(ws, row, col):
    value = ws.cell(row, col).value
    if value is None:
        return ""
    return str(value).strip()


def clean_text(value):
    value = re.sub(r"\s+", " ", value or "").strip()
    return value.replace("\ufffd", " ")


def normalize_theme(theme):
    theme = (theme or "").upper()
    theme = theme.split("(")[0].strip()
    theme = theme.replace(" ", "_")
    return theme


def fallback_attrs(title):
    attrs = {}
    clean = clean_text(title)
    m = re.search(r"\((\d{2,5}\s*/\s*\d{2,5})\s*Grit\)", clean, re.I)
    if not m:
        m = re.search(r"\((\d{2,5})\s*Grit\)", clean, re.I)
    if not m:
        m = re.search(r"\b(\d{2,5}\s*/\s*\d{2,5})\s*Grit\b", clean, re.I)
    if not m:
        m = re.search(r"\b(\d{2,5})\s*Grit\b", clean, re.I)
    if m:
        attrs["Style"] = re.sub(r"\s+", "", m.group(1)) + " Grit"

    m = re.search(r"\b(\d+)\s*pcs?\b", clean, re.I)
    if m:
        attrs["Number of Items"] = m.group(1)

    m = re.search(r"\((CBN|PCBN|PCD),\s*([^)]+)\)", clean, re.I)
    if m:
        attrs["Color"] = m.group(1).upper()
        attrs["Size"] = m.group(2).strip()

    return attrs


def choose_attrs(items):
    varied = []
    for name in ATTR_COLS:
        values = [str(item["raw_attrs"].get(name, "")).strip() for item in items]
        unique = {value for value in values if value}
        if len(unique) > 1:
            varied.append(name)

    theme = normalize_theme(items[0].get("variation_theme", ""))
    preferred = [name for name in THEME_ATTRS.get(theme, []) if name in varied]
    if preferred:
        names = preferred
    elif "Style" in varied:
        names = ["Style"]
    elif "Number of Items" in varied:
        names = ["Number of Items"]
    elif "Color" in varied and "Size" in varied:
        names = ["Color", "Size"]
    elif varied:
        names = varied[:2]
    else:
        names = []

    combos = set()
    has_duplicate = False
    for item in items:
        attrs = {name: item["raw_attrs"].get(name, "") for name in names if item["raw_attrs"].get(name, "") != ""}
        if not attrs:
            attrs = fallback_attrs(item["item_name"] or item["title"])
        if names:
            attrs = {name: attrs.get(name, "") for name in names if attrs.get(name, "") != ""}
        item["attrs"] = attrs
        combo = tuple((name, str(attrs.get(name, ""))) for name in sorted(attrs))
        if not combo or combo in combos:
            has_duplicate = True
        combos.add(combo)

    if has_duplicate:
        all_names = list(dict.fromkeys(names + varied[:3]))
        combos = set()
        duplicate_after_expand = False
        for item in items:
            attrs = {name: item["raw_attrs"].get(name, "") for name in all_names if item["raw_attrs"].get(name, "") != ""}
            if not attrs:
                attrs = fallback_attrs(item["item_name"] or item["title"])
            item["attrs"] = attrs
            combo = tuple((name, str(attrs.get(name, ""))) for name in sorted(attrs))
            if not combo or combo in combos:
                duplicate_after_expand = True
            combos.add(combo)
        has_duplicate = duplicate_after_expand

    return has_duplicate


def main():
    if len(sys.argv) != 3:
        print("Usage: extract-amazon-parent-child-map.py SOURCE.xlsm OUT.csv", file=sys.stderr)
        return 2

    source = Path(sys.argv[1])
    out = Path(sys.argv[2])
    wb = load_workbook(source, read_only=True, data_only=True, keep_vba=True)
    ws = wb["Template"]

    parents = {}
    groups = defaultdict(list)
    for row in range(7, ws.max_row + 1):
        sku = cell(ws, row, BASE_COLS["sku"])
        if not sku:
            continue
        record = {name: cell(ws, row, col) for name, col in BASE_COLS.items()}
        record["excel_row"] = row
        record["raw_attrs"] = {name: cell(ws, row, col) for name, col in ATTR_COLS.items() if cell(ws, row, col)}

        parentage = record["parentage"].lower()
        if parentage == "parent":
            parents[sku] = record
        elif parentage == "child" and record["parent_sku"]:
            groups[record["parent_sku"]].append(record)

    out.parent.mkdir(parents=True, exist_ok=True)
    with out.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=[
                "parent_sku",
                "parent_title",
                "variation_theme",
                "child_sku",
                "child_title",
                "child_status",
                "excel_row",
                "attrs_json",
                "duplicate_attrs",
            ],
        )
        writer.writeheader()
        written = 0
        duplicate_groups = 0
        for parent_sku, items in sorted(groups.items()):
            parent = parents.get(parent_sku, {})
            duplicate_attrs = choose_attrs(items)
            if duplicate_attrs:
                duplicate_groups += 1
            parent_title = clean_text(parent.get("item_name") or parent.get("title") or items[0].get("item_name") or items[0].get("title"))
            for item in items:
                writer.writerow({
                    "parent_sku": parent_sku,
                    "parent_title": parent_title,
                    "variation_theme": item.get("variation_theme", ""),
                    "child_sku": item["sku"],
                    "child_title": clean_text(item.get("item_name") or item.get("title")),
                    "child_status": item.get("status", ""),
                    "excel_row": item["excel_row"],
                    "attrs_json": json.dumps(item.get("attrs", {}), ensure_ascii=False, sort_keys=True),
                    "duplicate_attrs": "1" if duplicate_attrs else "0",
                })
                written += 1

    print(f"Parent groups: {len(groups)}")
    print(f"Child rows written: {written}")
    print(f"Groups with duplicate/weak attrs: {duplicate_groups}")


if __name__ == "__main__":
    raise SystemExit(main())
